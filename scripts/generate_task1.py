import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from doc_helper import save_and_archive

def apply_table_styles(table, dxa_widths=None):
    """
    Applies exact department styling:
    - Fixed table width: 7900 dxa (5.486 inches) fitting inside the 5.71-inch cell
    - No center misalignment (flush with text)
    - Fixed table layout (tblLayout type='fixed')
    - Clean single borders on all edges and inside cells (#CCCCCC)
    - Cell margins (padding): top=80, bottom=80, left=80, right=80 dxa
    - tblHeader on row 0 so headers repeat cleanly across page breaks
    - cantSplit on every row so rows don't get sliced across page breaks
    - Exact dxa column widths applied to every cell
    """
    tblPr = table._tbl.tblPr
    
    for jc in tblPr.xpath('./w:jc'):
        tblPr.remove(jc)
        
    for w in tblPr.xpath('./w:tblW'):
        tblPr.remove(w)
    tblPr.append(parse_xml(r'<w:tblW %s w:w="7900" w:type="dxa"/>' % nsdecls('w')))
    
    for l in tblPr.xpath('./w:tblLayout'):
        tblPr.remove(l)
    tblPr.append(parse_xml(r'<w:tblLayout %s w:type="fixed"/>' % nsdecls('w')))
    
    for b in tblPr.xpath('./w:tblBorders'):
        tblPr.remove(b)
    tblBorders = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    tblPr.append(tblBorders)
    
    for m in tblPr.xpath('./w:tblCellMar'):
        tblPr.remove(m)
    tblCellMar = parse_xml(r'''
        <w:tblCellMar %s>
            <w:top w:w="80" w:type="dxa"/>
            <w:bottom w:w="80" w:type="dxa"/>
            <w:left w:w="80" w:type="dxa"/>
            <w:right w:w="80" w:type="dxa"/>
        </w:tblCellMar>
    ''' % nsdecls('w'))
    tblPr.append(tblCellMar)

    header_trPr = table.rows[0]._tr.get_or_add_trPr()
    if not header_trPr.xpath('./w:tblHeader'):
        header_trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    if not header_trPr.xpath('./w:cantSplit'):
        header_trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
    
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        if not trPr.xpath('./w:cantSplit'):
            trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        if dxa_widths:
            for idx, dxa_w in enumerate(dxa_widths):
                if idx < len(row.cells):
                    cell = row.cells[idx]
                    tcPr = cell._tc.get_or_add_tcPr()
                    for w in tcPr.xpath('./w:tcW'):
                        tcPr.remove(w)
                    tcPr.append(parse_xml(r'<w:tcW %s w:w="%s" w:type="dxa"/>' % (nsdecls('w'), str(dxa_w))))

def style_cell(cell, bg_color=None, bold=False, font_size=9.5, font_name="Calibri", text=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    if bg_color:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)
    if text is not None:
        cell.text = text
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs:
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.bold = bold

def add_heading(cell, text, bold=True, size=12, keep_with_next=True):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep_with_next
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    return p

def add_body_p(cell, text, bold_prefix=None, space_after=3):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(11)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    return p

def add_figure(cell, image_path, caption_text, width_inches=5.4):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(image_path, width=Inches(width_inches))
    
    cap = cell.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_with_next = True
    c_run = cap.add_run(caption_text)
    c_run.bold = True
    c_run.font.name = "Calibri"
    c_run.font.size = Pt(10)
    return p, cap

def generate_task1_document():
    base_dir = r"d:\Devops_project"
    template_path = os.path.join(base_dir, "templates", "template_task1.docx")
    
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Base template not found at {template_path}")
        
    doc = docx.Document(template_path)
    
    # Clean Table 0 encoding issues
    t0 = doc.tables[0]
    t0.rows[0].cells[0].text = "Semester"
    t0.rows[0].cells[1].text = "B.E. Semester VII - Computer Engineering"
    t0.rows[1].cells[0].text = "Subject"
    t0.rows[1].cells[1].text = "DevOps Lab"
    t0.rows[2].cells[0].text = "Subject In-charge"
    t0.rows[2].cells[1].text = "Dr. Snehal Andhare"
    for row in t0.rows:
        style_cell(row.cells[0], bold=True, font_size=10.5)
        style_cell(row.cells[1], bold=False, font_size=10.5)
        
    # Clean Table 1 encoding issues
    t1 = doc.tables[1]
    t1.rows[0].cells[0].text = "Student Name"
    t1.rows[0].cells[1].text = "Rushikesh Gawade"
    t1.rows[1].cells[0].text = "Roll Number"
    t1.rows[1].cells[1].text = "23102A0052"
    t1.rows[2].cells[0].text = "Grade and Subject \nTeacher's Signature"
    for r_idx in range(2):
        style_cell(t1.rows[r_idx].cells[0], bold=True, font_size=10.5)
        style_cell(t1.rows[r_idx].cells[1], bold=False, font_size=10.5)
    style_cell(t1.rows[2].cells[0], bold=True, font_size=10.0)

    # Table 2: 4-row layout for Task 1
    t2 = doc.tables[2]
    
    # Row 0: Task Number
    t2.rows[0].cells[0].text = "Task Number"
    style_cell(t2.rows[0].cells[0], bold=True, font_size=11)
    t2.rows[0].cells[1].text = "1"
    style_cell(t2.rows[0].cells[1], bold=True, font_size=12, text="1")
    
    # Row 1: Task Title
    t2.rows[1].cells[0].text = "Task Title"
    style_cell(t2.rows[1].cells[0], bold=True, font_size=11)
    title_text = "Problem Definition and Scope: Study the real-time need for the Customer Review Moderation Portal. Identify target users, existing pain points, stakeholders, constraints and measurable success criteria. Freeze a small 15-Task MVP scope."
    t2.rows[1].cells[1].text = title_text
    style_cell(t2.rows[1].cells[1], bold=True, font_size=10.5, text=title_text)
    
    # Row 2: Program & Output (Core content cell)
    t2.rows[2].cells[0].text = "Program & Output"
    style_cell(t2.rows[2].cells[0], bold=True, font_size=11)
    
    body_cell = t2.rows[2].cells[1]
    if body_cell.paragraphs:
        body_cell.paragraphs[0].text = ""
        p0 = body_cell.paragraphs[0]
    else:
        p0 = body_cell.add_paragraph()
        
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.keep_with_next = True
    r0 = p0.add_run("Executive summary and real-time need analysis")
    r0.bold = True
    r0.font.name = "Calibri"
    r0.font.size = Pt(12)
    
    add_body_p(body_cell, "In modern e-commerce platforms, software-as-a-service (SaaS) ecosystems, and consumer marketplaces, user-generated content (UGC)—specifically customer product reviews and star ratings—serves as the primary pillar of buyer trust, brand reputation, and algorithmic product discoverability. Empirical market research indicates that over 93% of prospective consumers consult online reviews prior to making purchasing decisions, and a single-star improvement in aggregate product rating can drive a 5% to 9% increase in commercial revenue. Consequently, the customer review stream has become a prime target for adversarial manipulation, malicious review bombing, automated spam bots, abusive commentary, and fraudulent competitor sabotage.")
    add_body_p(body_cell, "Operating an e-commerce platform without an automated, real-time Customer Review Moderation Portal (CRMP) introduces severe business vulnerabilities. Unmoderated systems risk publishing toxic or defamatory content, violating digital compliance mandates (such as the Federal Trade Commission's guidelines against deceptive endorsements), and eroding customer confidence. Conversely, purely manual moderation workflows in legacy enterprises rely on fragmented spreadsheets or disconnected administrative dashboards, leading to operational backlogs where genuine reviews remain unpublished for days. To resolve these challenges, I have undertaken the engineering and DevOps lifecycle implementation of the Customer Review Moderation Portal.")
    add_body_p(body_cell, "From an engineering perspective, this platform must be architected with deterministic reliability, zero external database setup overhead, and lightning-fast local and CI/CD execution. To achieve seamless continuous integration and deployment across all 15 tasks of our university laboratory syllabus, the portal is developed using Java 21 LTS with Spring Boot 3.x, an embedded in-memory H2 database, Thymeleaf responsive server-side UI with Bootstrap 5, and Spring Boot Actuator telemetry for automated liveness probing. This ensures that every git commit can be compiled, tested via headless Selenium WebDriver, packaged into lightweight multi-stage Docker containers, and provisioned idempotently using Ansible automation.")

    # Section 1: Detailed Problem Statement
    add_heading(body_cell, "Detailed problem statement: Functional and DevOps challenges", keep_with_next=True)
    add_body_p(body_cell, "The engineering design of the Customer Review Moderation Portal addresses two tightly coupled problem domains: operational review governance challenges and software delivery lifecycle inefficiencies:")
    add_body_p(body_cell, "1. Operational and Content Moderation Bottlenecks: Customer feedback arrives continuously across diverse product categories. Without automated validation gates and structured workflow state machines, moderators face overwhelming queues. Content moderators lack unified tools to inspect review sentiment, evaluate star ratings against comment text, flag suspected spam, and make binding decisions (Approve, Reject, Flag). Furthermore, legacy systems lack strict role separation, enabling accidental deletion or unvetted approval of abusive remarks.", bold_prefix="• ")
    add_body_p(body_cell, "2. DevOps and Continuous Delivery Inefficiencies: Traditional review systems suffer from brittle deployment cycles. Code changes to moderation business rules are frequently deployed manually to production servers without automated regression testing, causing unexpected downtime or broken UI locators. Inconsistent environment configurations between developer workstations, staging nodes, and production servers produce 'it works on my machine' anomalies. Without containerization and infrastructure-as-code automation, disaster recovery and rolling deployments remain high-risk manual interventions.", bold_prefix="• ")

    # Figure 1.1: System Context & Review Moderation Process Architecture
    fig1_path = os.path.join(base_dir, "screenshots", "task1", "figure_1_1_context_architecture.png")
    if os.path.exists(fig1_path):
        add_figure(body_cell, fig1_path, 
                   "Figure 1.1: Architectural Context and Review Moderation Workflow of the Customer Review Moderation Portal", 
                   width_inches=5.4)

    # Section 2: Target Users and Stakeholder Matrix
    add_heading(body_cell, "Target users and stakeholder governance matrix", keep_with_next=True)
    add_body_p(body_cell, "To ensure comprehensive governance across the application and its operational lifecycle, I conducted a formal stakeholder analysis. The portal serves both business-facing actors and engineering operations personnel:")

    # Table 1: Stakeholder Identification & RACI Matrix (Total = 7900 dxa: 1100, 1500, 2400, 2000, 900)
    table_sh = body_cell.add_table(rows=7, cols=5)
    apply_table_styles(table_sh, dxa_widths=[1100, 1500, 2400, 2000, 900])
    
    headers_sh = ["Stakeholder Role", "Category", "Operational Responsibility", "Key DevOps / System Concern", "RACI"]
    for i, h in enumerate(headers_sh):
        style_cell(table_sh.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    sh_data = [
        ("End Shopper / Customer", "External Consumer", "Submits product reviews, ratings (1–5 stars), titles, and feedback commentary post-purchase.", "Intuitive UI, input error feedback, rapid review publication without data loss.", "Informed (I)"),
        ("Content Moderator", "Internal Operations", "Inspects pending moderation queue, filters reviews, and approves, rejects, or flags submissions.", "Low query latency (< 250ms), reliable search filters, dynamic status tag updates.", "Responsible (R)"),
        ("Compliance Officer / Lead", "Governance & Legal", "Audits flagged suspicious reviews, enforces anti-spam and anti-defamation consumer laws.", "Immutable audit logs, git commit traceability, zero unauthorized status overrides.", "Accountable (A)"),
        ("Product Merchant / Brand Owner", "Business Client", "Monitors aggregated customer sentiment, average product ratings, and feedback volume via KPI cards.", "Real-time dashboard accuracy, zero cache inconsistency on approved counts.", "Consulted (C)"),
        ("DevOps Platform Engineer", "Technical Engineering", "Maintains Git branching, Jenkins CI/CD pipelines, Selenium automated tests, Docker images, and Ansible scripts.", "100% reproducible builds, automated failure gates, idempotency, zero-downtime rollouts.", "Responsible (R)"),
        ("Academic Evaluator (Faculty)", "University Examination", "Evaluates student implementation against official B.E. Semester VII DevOps syllabus and criteria.", "Strict adherence to academic identity, zero mock placeholders, reproducible live execution.", "Accountable (A)")
    ]
    for r_idx, row_vals in enumerate(sh_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_sh.rows[r_idx].cells[c_idx], font_size=8.8, text=val)

    # Section 3: Existing Pain Points vs Target Solution
    add_heading(body_cell, "Analysis of existing pain points and workflow inefficiencies", keep_with_next=True)
    add_body_p(body_cell, "A rigorous comparison between existing legacy review moderation mechanisms and our target Git-based DevOps moderation architecture illustrates the transformative impact of automated pipelines:")

    # Table 2: Existing Pain Points vs DevOps Solution (Total = 7900 dxa: 1300, 2200, 2600, 1800)
    table_pp = body_cell.add_table(rows=7, cols=4)
    apply_table_styles(table_pp, dxa_widths=[1300, 2200, 2600, 1800])
    
    headers_pp = ["Dimension", "Legacy / Manual Approach", "Target Git-Based DevOps Portal", "Expected Engineering Impact"]
    for i, h in enumerate(headers_pp):
        style_cell(table_pp.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    pp_data = [
        ("Review Ingestion & Validation", "Client-side validation easily bypassed; unvalidated inputs crash backend or corrupt data stores.", "Robust Spring Boot Bean Validation (@NotNull, @Min(1), @Max(5), @Size) with inline UI alerts.", "Eliminates bad payloads and malformed data at the application threshold."),
        ("Moderation Queue Triage", "Moderators review raw rows in shared spreadsheets; high collision rate and status confusion.", "Interactive web directory with keyword search and status filters (PENDING, APPROVED, REJECTED, FLAGGED).", "Reduces average review triage latency from 48 hours to under 30 minutes."),
        ("Status Lifecycle Enforcement", "Status values updated manually with no transition rules; rejected items can accidentally publish.", "Finite state machine enforcing strict transitions: PENDING -> APPROVED / REJECTED / FLAGGED.", "Guarantees zero accidental data leaks of toxic or flagged content."),
        ("Quality Assurance & Testing", "Manual ad-hoc testing before deployment; regression bugs undetected until production failure.", "Headless Selenium WebDriver test suite covering 4 critical user journeys with failure screenshots.", "100% automated regression gate in CI; prevents defect leakage to production."),
        ("Build & Deployment Automation", "Manual FTP/SSH file transfer to runtime servers; high configuration drift and human error.", "Jenkins continuous integration and pipeline-as-code (Jenkinsfile) triggered automatically.", "Deterministic, repeatable builds completed in under 4 minutes with zero manual intervention."),
        ("Infrastructure Management", "Manual server configuration; undocumented package installs leading to configuration drift.", "Ansible automated playbooks enforcing exact prerequisites, directory permissions, and firewall rules.", "Guarantees 100% idempotent infrastructure (changed=0 on re-run) and rapid disaster rollback.")
    ]
    for r_idx, row_vals in enumerate(pp_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_pp.rows[r_idx].cells[c_idx], font_size=8.8, text=val)

    # Section 4: SMART Objectives
    add_heading(body_cell, "Engineering and operational objectives (SMART criteria)", keep_with_next=True)
    add_body_p(body_cell, "To ensure that project delivery remains focused, quantifiable, and academically rigorous, I formulated the following objectives using the SMART (Specific, Measurable, Achievable, Relevant, Time-bound) framework:")
    add_body_p(body_cell, "Specific: Design, develop, containerize, and deploy a responsive Customer Review Moderation Portal featuring review ingestion, search filtering, role-based status lifecycle, a real-time KPI metrics dashboard, and an Actuator health check endpoint.", bold_prefix="1. ")
    add_body_p(body_cell, "Measurable: Achieve a 100% pass rate across headless Selenium WebDriver test suites, ensure zero compiler warnings, maintain an application response latency under 300 ms, and verify that Ansible automation achieves idempotency (changed=0, failed=0).", bold_prefix="2. ")
    add_body_p(body_cell, "Achievable: Implement the system utilizing proven enterprise open-source technologies—Java 21 LTS, Spring Boot 3.x, Embedded H2 Database, Apache Maven Wrapper, Docker, and Jenkins—eliminating external licensing or cloud infrastructure barriers.", bold_prefix="3. ")
    add_body_p(body_cell, "Relevant: Direct alignment with the university B.E. Computer Engineering DevOps curriculum, systematically exercising version control, continuous integration, continuous testing, container lifecycle, and automated provisioning.", bold_prefix="4. ")
    add_body_p(body_cell, "Time-Bound: Execute and document the complete 15-task syllabus progression within the allocated 15 laboratory milestones, ensuring every artifact is verified, timestamped, and cataloged deterministically.", bold_prefix="5. ")

    # Section 5: Measurable Success Criteria & KPIs
    add_heading(body_cell, "Measurable success criteria, SLAs, and DevOps quality metrics", keep_with_next=True)
    add_body_p(body_cell, "The success of the Customer Review Moderation Portal is verified against predefined, quantitatively measurable Service Level Agreements (SLAs) and DevOps quality benchmarks:")

    # Table 3: Success Criteria & SLAs (Total = 7900 dxa: 900, 1800, 1500, 1800, 1900)
    table_sc = body_cell.add_table(rows=7, cols=5)
    apply_table_styles(table_sc, dxa_widths=[900, 1800, 1500, 1800, 1900])
    
    headers_sc = ["Metric ID", "Quality Dimension", "Legacy Baseline", "Target SLA in DevOps Portal", "Verification / Measurement Method"]
    for i, h in enumerate(headers_sc):
        style_cell(table_sc.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    sc_data = [
        ("SLA-01", "Review Submission Latency", "1200 ms - 2500 ms", "< 250 ms local; < 500 ms under load", "Browser DevTools network timing & cURL probes"),
        ("SLA-02", "Automated Test Pass Rate", "0% (Manual testing)", "100% across all 4 critical user journeys", "Headless Selenium WebDriver 4.x test execution"),
        ("SLA-03", "CI Quality Gate Halt", "Manual code reviews only", "Immediate build abortion on test failure", "Jenkins Surefire test quality gate assertion"),
        ("SLA-04", "Docker Build & Packaging", "Manual jar export (>15m)", "< 90 seconds multi-stage image build", "Docker CLI build timing & Temurin JRE runtime"),
        ("SLA-05", "Configuration Idempotency", "Unchecked configuration drift", "Exact idempotency (changed=0, failed=0)", "Ansible playbook re-execution log audit"),
        ("SLA-06", "Service Liveness Verification", "Manual page inspection", "HTTP 200 OK with {\"status\":\"UP\"}", "Automated curl check against /actuator/health")
    ]
    for r_idx, row_vals in enumerate(sc_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_sc.rows[r_idx].cells[c_idx], font_size=8.8, text=val)

    # Section 6: Technical and Operational Constraints
    add_heading(body_cell, "Technical, architectural, and operational constraints", keep_with_next=True)
    add_body_p(body_cell, "The engineering design respects strict operational boundaries to guarantee consistent execution across diverse university evaluation environments:")

    # Table 4: System Constraints (Total = 7900 dxa: 1300, 2100, 2500, 2000)
    table_co = body_cell.add_table(rows=6, cols=4)
    apply_table_styles(table_co, dxa_widths=[1300, 2100, 2500, 2000])
    
    headers_co = ["Constraint Category", "Constraint Description", "Operational Boundary & Impact", "Engineering Mitigation"]
    for i, h in enumerate(headers_co):
        style_cell(table_co.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    co_data = [
        ("Runtime & Database", "Zero external database daemon dependencies", "Precludes external MySQL/PostgreSQL setup that may fail during automated CI test runs.", "Embedded in-memory H2 database with Spring Data JPA; resets cleanly per test run."),
        ("Port Allocation", "Single designated host port binding", "Host port 8080 must be mapped consistently across local, Docker, and Ansible nodes.", "Configured server.port=8080 with environment variable overrides for container mapping."),
        ("Container Footprint", "Lightweight container image size constraint", "Bulky full-JDK images (> 600 MB) consume excessive disk and slow down CI pipeline staging.", "Multi-stage Dockerfile using minimal Eclipse Temurin JRE runtime (< 220 MB total image)."),
        ("UI Automation", "Headless browser execution compatibility", "Automated CI servers (Jenkins) lack physical graphics displays and X-servers.", "Configured Chrome/Edge headless execution flags (--headless, --no-sandbox, --disable-dev-shm-usage)."),
        ("Version Control", "Protected main branch with strict PR review", "Prevents unvetted or broken code from bypassing the integration and testing gates.", "Feature branch policy (feature/* -> dev -> main) with mandatory pull request merges.")
    ]
    for r_idx, row_vals in enumerate(co_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_co.rows[r_idx].cells[c_idx], font_size=8.8, text=val)

    # Section 7: Approved 15-Task MVP Scope Breakdown
    add_heading(body_cell, "Approved 15-Task MVP scope and curriculum syllabus roadmap", keep_with_next=True)
    add_body_p(body_cell, "To establish complete alignment between the university DevOps curriculum and the Customer Review Moderation Portal, I decomposed the semester project into 15 structured, sequential milestones. Each milestone delivers a verified, production-ready engineering artifact:")

    # Table 5: 15-Task Scope Breakdown (Total = 7900 dxa: 650, 1600, 2450, 1800, 1400)
    table_mvp = body_cell.add_table(rows=16, cols=5)
    apply_table_styles(table_mvp, dxa_widths=[650, 1600, 2450, 1800, 1400])
    
    headers_mvp = ["Task", "Syllabus Task Title", "Engineering Scope & Activity", "DevOps Deliverable", "Quality Gate / Verification"]
    for i, h in enumerate(headers_mvp):
        style_cell(table_mvp.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    mvp_data = [
        ("Task 1", "Problem Definition and Scope", "Analyze real-time moderation needs, identify stakeholders, pain points, constraints, and freeze 15-task scope.", "Problem statement, stakeholder matrix, MVP scope document.", "Academic approval and syllabus alignment review."),
        ("Task 2", "Agile Planning & DevOps Workflow", "Author user stories, acceptance criteria, Scrum/Kanban backlog, Definition of Done, and lifecycle diagram.", "Product backlog, task board, sprint plan, DevOps flow diagram.", "Complete acceptance criteria validation."),
        ("Task 3", "Requirements, Architecture & Setup", "Architect 4-tier Spring Boot app, SRS summary, use-case model, H2 database schema, and working local baseline.", "SRS summary, use-case diagram, data models, local baseline.", "Local compilation and Actuator health check (UP)."),
        ("Task 4", "Git & GitHub Initialization", "Initialize Git repository, configure .gitignore, issue templates, branch policies, and initial skeleton commit.", "Repository URL, README, .gitignore, branch policy, commits.", "Clean commit history adhering to Conventional Commits."),
        ("Task 5", "Feature Dev with Branching", "Implement review ingestion workflow on feature branch, demonstrate git add/commit/push/PR/merge workflow.", "Feature branch, working feature 1, pull request, merge evidence.", "Successful peer review and fast-forward/merge commit."),
        ("Task 6", "MVP Completion & Collaboration", "Complete remaining 4 MVP modules, engineer intentional merge conflict, resolve manually, and tag v1.0.0.", "Functional MVP, resolved conflict log, tagged release v1.0.0.", "Zero merge conflicts; release baseline verified."),
        ("Task 7", "Jenkins CI Job Configuration", "Install Jenkins, connect GitHub repo, create automated Maven build job with commit polling trigger.", "Configured Jenkins CI job, build logs, archived .jar artifact.", "Automated trigger on commit; BUILD SUCCESS."),
        ("Task 8", "Pipeline as Code & Deployment", "Author multi-stage Jenkinsfile (Checkout, Build, Test, Deploy) with parameterized deployment targets.", "Jenkinsfile, declarative pipeline execution, deployment URL.", "End-to-end pipeline execution with zero manual steps."),
        ("Task 9", "Selenium Test Design & Execution", "Design Selenium WebDriver suite covering 4 critical user journeys with assertions and failure screenshot hooks.", "Selenium test plan, Java test classes, failure screenshot logic.", "100% headless pass rate via mvn test."),
        ("Task 10", "Continuous Testing in Jenkins", "Integrate Selenium tests into Jenkins pipeline; demonstrate failure gate on intentional bug and subsequent fix.", "Jenkins test report, aborted pipeline evidence, fix commit run.", "Pipeline immediately halts on defect assertion failure."),
        ("Task 11", "Docker Image & Container Lifecycle", "Author multi-stage Dockerfile, build/tag image, map port 8080, inspect logs, and manage container lifecycle.", "Dockerfile, image inspection, container lifecycle command logs.", "Container runs cleanly in isolation on port 8080."),
        ("Task 12", "Jenkins-Docker Continuous CD", "Extend Jenkinsfile to automatically build versioned Docker image, publish to registry, and run new container.", "Versioned container image, automated CD pipeline execution.", "Automated commit-to-running-container deployment."),
        ("Task 13", "Configuration Management Script", "Identify server prerequisites (Java, Docker, directories, ports) and author structured Ansible playbook.", "Ansible inventory (hosts.ini), playbook.yml, execution log.", "Syntax validation via ansible-playbook --syntax-check."),
        ("Task 14", "Automated Provisioning & Reliability", "Provision clean target node, verify idempotency (changed=0), test /actuator/health, and execute rollback.", "Provisioned host, idempotency evidence, rollback log.", "Zero changes on repeated execution; health verified."),
        ("Task 15", "Final End-to-End Release & Viva", "Execute complete commit-to-production pipeline, assemble final comprehensive report, prepare viva demonstration.", "Production release, end-to-end audit, complete report, viva.", "Comprehensive defense and faculty evaluation.")
    ]
    for r_idx, row_vals in enumerate(mvp_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_mvp.rows[r_idx].cells[c_idx], font_size=8.5, text=val)

    # Figure 1.2: 15-Task End-to-End DevOps Lifecycle Roadmap
    fig2_path = os.path.join(base_dir, "screenshots", "task1", "figure_1_2_devops_roadmap.png")
    if os.path.exists(fig2_path):
        add_figure(body_cell, fig2_path, 
                   "Figure 1.2: 15-Task End-to-End DevOps Continuous Lifecycle and Quality Delivery Roadmap", 
                   width_inches=5.4)

    # Section 8: Core MVP Functional Modules
    add_heading(body_cell, "Core MVP functional scope and state machine architecture", keep_with_next=True)
    add_body_p(body_cell, "To ensure rapid execution, zero architectural bloat, and highly focused Selenium test locator targeting, the MVP scope of the Customer Review Moderation Portal is strictly frozen to five core modules:")
    add_body_p(body_cell, "1. Review Ingestion & Input Validation: A responsive submission interface where shoppers submit Product ID, Reviewer Name, Rating (1–5 Stars), Review Title, and Feedback Comments. Backend Spring Bean Validation enforces mandatory constraints and rejects empty or out-of-bound payloads with clear inline visual feedback.", bold_prefix="• ")
    add_body_p(body_cell, "2. Review Directory & Moderation Queue: A centralized moderation table providing real-time text search across product names and review comments, accompanied by status filtering (All, Pending, Approved, Rejected, Flagged) to triage backlogs efficiently.", bold_prefix="• ")
    add_body_p(body_cell, "3. Role-Based Moderation Status Lifecycle: A deterministic state machine governing review progression from initial ingestion to public publication or quarantine, strictly disallowing illegal state transitions.", bold_prefix="• ")
    add_body_p(body_cell, "4. Summary Metrics & KPI Analytics Dashboard: Real-time calculation cards rendering Total Reviews, Pending Queue Count, Approved Reviews, Flagged/Rejected Count, and Average Aggregate Star Rating, providing instant operational visibility.", bold_prefix="• ")
    add_body_p(body_cell, "5. Spring Boot Actuator Telemetry Endpoint: Standardized HTTP GET /actuator/health endpoint exposing JSON liveness/readiness telemetry (status: UP) for automated Jenkins, Docker, and Ansible health checks.", bold_prefix="• ")

    # Table 6: State Machine Transitions (Total = 7900 dxa: 1200, 1400, 1200, 1200, 2900)
    table_sm = body_cell.add_table(rows=6, cols=5)
    apply_table_styles(table_sm, dxa_widths=[1200, 1400, 1200, 1200, 2900])
    
    headers_sm = ["Current State", "Trigger Event", "Permitted Role", "Next State", "Guard Conditions & Business Action"]
    for i, h in enumerate(headers_sm):
        style_cell(table_sm.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    sm_data = [
        ("NONE", "Submit Review", "End Shopper", "PENDING", "All fields valid; rating between 1 and 5; comment > 5 chars; persisted to H2 queue."),
        ("PENDING", "Approve Review", "Content Moderator", "APPROVED", "Moderator confirms content complies with community guidelines; review becomes publicly visible."),
        ("PENDING", "Reject Review", "Content Moderator", "REJECTED", "Content contains spam, profanity, or competitor slander; suppressed from public view."),
        ("PENDING", "Flag for Audit", "Content Moderator", "FLAGGED", "Ambiguous feedback or suspected fraudulent pattern; escalated to Senior Compliance Officer."),
        ("FLAGGED", "Compliance Resolve", "Compliance Lead", "APPROVED / REJECTED", "Senior audit confirms validity or confirms violation; final terminal state assigned.")
    ]
    for r_idx, row_vals in enumerate(sm_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_sm.rows[r_idx].cells[c_idx], font_size=8.8, text=val)

    # Section 9: Integrated DevOps Toolchain Ecosystem
    add_heading(body_cell, "Integrated DevOps toolchain and deployment architecture", keep_with_next=True)
    add_body_p(body_cell, "To realize continuous quality and automated deployment across all 15 tasks, the engineering ecosystem integrates industry-standard DevOps tools:")

    # Table 7: DevOps Toolchain (Total = 7900 dxa: 1300, 1700, 1900, 3000)
    table_tc = body_cell.add_table(rows=9, cols=4)
    apply_table_styles(table_tc, dxa_widths=[1300, 1700, 1900, 3000])
    
    headers_tc = ["DevOps Lifecycle Stage", "Selected Technology", "Version / Runtime", "Role & Engineering Responsibility in Portal"]
    for i, h in enumerate(headers_tc):
        style_cell(table_tc.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    tc_data = [
        ("Source Code Management", "Git & GitHub", "Git 2.51+ (Windows/Linux)", "Distributed version control, branch protection (main/dev), pull requests, and commit traceability."),
        ("Build & Dependency Tool", "Apache Maven Wrapper", "Maven 3.9.x / mvnw", "Deterministic compilation, automated dependency resolution, and portable execution without pre-installed Maven."),
        ("Programming Platform", "Java & Spring Boot", "Java 21 LTS / Spring Boot 3.x", "Enterprise backend web application framework providing Spring MVC, Data JPA, and Actuator modules."),
        ("In-Memory Persistence", "H2 Embedded Database", "H2 Database Engine", "Zero-configuration in-memory persistence ensuring lightning-fast test execution and reproducible CI builds."),
        ("Continuous Integration", "Jenkins Automation Server", "Jenkins LTS 2.440+", "Automated build triggers, pipeline-as-code (Jenkinsfile) orchestration, and test result archiving."),
        ("Continuous Testing Gate", "Selenium WebDriver & JUnit", "Selenium 4.x / JUnit 5", "Headless browser end-to-end user journey validation with automated failure screenshot capture."),
        ("Containerization & CD", "Docker Engine", "Docker 24.x+ / Temurin JRE", "Lightweight multi-stage container packaging for portable deployment on port 8080."),
        ("Configuration Management", "Ansible Automation", "Ansible Core 2.15+", "Idempotent host provisioning, environment configuration, dependency installation, and rollback.")
    ]
    for r_idx, row_vals in enumerate(tc_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_tc.rows[r_idx].cells[c_idx], font_size=8.8, text=val)

    # Row 3: Conclusion (Academic Summary)
    t2.rows[3].cells[0].text = "Conclusion:"
    style_cell(t2.rows[3].cells[0], bold=True, font_size=11)
    
    conc_text = "In this initial task, I have successfully formulated the comprehensive problem definition, analyzed operational real-time requirements, and established the boundary constraints for the Git-Based Customer Review Moderation Portal (CRMP). By conducting an in-depth study of e-commerce user feedback ecosystems, I identified the critical risks associated with unmoderated content—such as spam injection, competitor review bombing, and regulatory non-compliance—as well as the software delivery bottlenecks stemming from manual deployment processes and lack of automated quality gates. I established a structured stakeholder governance matrix with RACI assignments, defined measurable Service Level Agreements (SLAs) including sub-500ms review ingestion latency and 100% automated test pass rates, and articulated multi-dimensional system constraints that guarantee zero external database overhead via an embedded in-memory H2 database. Furthermore, I systematically froze the 15-Task MVP scope, aligning every syllabus milestone from Agile backlog creation and Git collaboration to Jenkins CI pipelines, headless Selenium WebDriver gates, Docker containerization, and idempotent Ansible configuration management. The resulting foundational blueprint ensures a deterministic, highly automated, and academically rigorous platform ready for agile planning and full engineering realization in Task 2."
    t2.rows[3].cells[1].text = conc_text
    style_cell(t2.rows[3].cells[1], font_size=10.5, text=conc_text)

    output_path = os.path.join(base_dir, "docs", "DevOps_Task1_RushikeshGawade_23102A0052.docx")
    save_and_archive(doc, output_path)

if __name__ == "__main__":
    generate_task1_document()
